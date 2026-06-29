"""
Verifier for payment-fee-audit (4-rule fee calculation).

Rules from PDF:
1. Base rate by card type + amount tier
2. Dispute surcharge: 2x if is_disputed=1
3. Volume discount: 10% off if merchant monthly volume > $5000
4. Fee cap: max $25.00

Expected:
  Total final fees: 852.7417
  Merchants with volume discount: 5 merchants show YES

A model that misses the dispute surcharge, volume discount, or fee cap
will produce wrong totals.
"""
from openpyxl import load_workbook
from pathlib import Path

OUTPUT_FILE = Path("/root/fee_audit.xlsx")
TOL = 1.0  # $1 tolerance on totals


def header_map(ws, row=1):
    return {str(c.value).strip(): i for i, c in enumerate(ws[row], 1) if c.value}


class TestPaymentFeeAudit:
    def test_output_exists(self):
        assert OUTPUT_FILE.exists(), f"Missing: {OUTPUT_FILE}"

    def test_required_sheets(self):
        wb = load_workbook(OUTPUT_FILE, data_only=True)
        assert "Transaction Fees" in wb.sheetnames, f"Missing 'Transaction Fees'. Found: {wb.sheetnames}"
        assert "Merchant Summary" in wb.sheetnames, f"Missing 'Merchant Summary'. Found: {wb.sheetnames}"

    def test_transaction_fees_headers(self):
        wb = load_workbook(OUTPUT_FILE, data_only=True)
        cols = header_map(wb["Transaction Fees"])
        for h in ["transaction_id","card_type","amount","is_disputed","monthly_volume","base_fee","final_fee"]:
            assert h in cols, f"Missing column '{h}' in Transaction Fees"

    def test_row_count(self):
        wb = load_workbook(OUTPUT_FILE, data_only=True)
        rows = [r for r in wb["Transaction Fees"].iter_rows(min_row=2, values_only=True) if r[0]]
        assert len(rows) == 200, f"Expected 200 rows, got {len(rows)}"

    def test_total_fees_correct(self):
        """
        Total final_fee must be ~852.74.
        Missing dispute surcharge gives ~820. Missing volume discount gives ~870.
        Missing both gives ~838. Only all 4 rules give ~852.74.
        """
        wb = load_workbook(OUTPUT_FILE, data_only=True)
        ws = wb["Transaction Fees"]
        cols = header_map(ws)
        total = sum(float(r[cols["final_fee"]-1]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0])
        assert abs(total - 852.7417) < TOL, (
            f"Total final_fee={total:.4f}, expected ~852.7417. "
            f"Check: dispute surcharge (2x if is_disputed=1), "
            f"volume discount (10% off if merchant monthly_volume > $5000), "
            f"fee cap ($25.00 max)."
        )

    def test_volume_discount_applied_count(self):
        """5 merchants should have volume_discount_applied=YES."""
        wb = load_workbook(OUTPUT_FILE, data_only=True)
        ws = wb["Merchant Summary"]
        cols = header_map(ws)
        yes_count = sum(1 for r in ws.iter_rows(min_row=2, values_only=True)
                       if r[0] and str(r[cols["volume_discount_applied"]-1]) == "YES")
        assert yes_count == 5, f"Expected 5 merchants with YES, got {yes_count}"

    def test_merchant_summary_sorted(self):
        wb = load_workbook(OUTPUT_FILE, data_only=True)
        ws = wb["Merchant Summary"]
        cols = header_map(ws)
        ids = [int(r[cols["merchant_id"]-1]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
        assert ids == sorted(ids), f"Merchant Summary not sorted by merchant_id: {ids}"
