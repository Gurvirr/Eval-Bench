"""
Verifier for pdf-invoice-reconciliation (Excel output with strict format).

Agent must produce /root/reconciliation.xlsx with:
- Sheet "Ledger vs Invoices": exact headers, 3 rows sorted by invoice_ref
  - status: exactly "MATCH" or "DISCREPANCY" (case-sensitive)
- Sheet "Discrepancy Report": exact headers, type values must be UPPERCASE

Traps:
- Wrong sheet names
- Wrong status/type string values (case)
- Missing TechPro discrepancy (invoice with no ledger entries)
- Missing phantom entry (ledger with no invoice)
"""
from openpyxl import load_workbook
from pathlib import Path

OUTPUT_FILE = Path("/root/reconciliation.xlsx")
TOL = 0.05
VALID_TYPES = {"AMOUNT_MISMATCH","MISSING_FROM_LEDGER","PHANTOM_ENTRY"}


def header_map(ws, row=1):
    return {str(c.value).strip(): i for i, c in enumerate(ws[row], 1) if c.value}


class TestInvoiceReconciliation:
    def test_output_exists(self):
        assert OUTPUT_FILE.exists(), f"Missing: {OUTPUT_FILE}"

    def test_required_sheets(self):
        wb = load_workbook(OUTPUT_FILE, data_only=True)
        assert "Ledger vs Invoices" in wb.sheetnames, \
            f"Missing 'Ledger vs Invoices'. Found: {wb.sheetnames}"
        assert "Discrepancy Report" in wb.sheetnames, \
            f"Missing 'Discrepancy Report'. Found: {wb.sheetnames}"

    def test_ledger_vs_invoices_headers(self):
        wb = load_workbook(OUTPUT_FILE, data_only=True)
        cols = header_map(wb["Ledger vs Invoices"])
        for h in ["invoice_ref","vendor","invoice_total","ledger_total","difference","status"]:
            assert h in cols, f"Missing column '{h}' in Ledger vs Invoices"

    def test_three_invoice_rows(self):
        wb = load_workbook(OUTPUT_FILE, data_only=True)
        ws = wb["Ledger vs Invoices"]
        rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
        assert len(rows) == 3, f"Expected 3 invoice rows, got {len(rows)}"

    def test_techpro_marked_discrepancy(self):
        """INV-2024-002 has no ledger entries — must be DISCREPANCY."""
        wb = load_workbook(OUTPUT_FILE, data_only=True)
        ws = wb["Ledger vs Invoices"]
        cols = header_map(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == "INV-2024-002":
                status = str(row[cols["status"]-1])
                assert status == "DISCREPANCY", \
                    f"INV-2024-002 status='{status}', expected 'DISCREPANCY' (no ledger entries)"
                return
        assert False, "INV-2024-002 not found in sheet"

    def test_discrepancy_type_values(self):
        """type column must use exact UPPERCASE values."""
        wb = load_workbook(OUTPUT_FILE, data_only=True)
        ws = wb["Discrepancy Report"]
        cols = header_map(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]: continue
            t = str(row[cols["type"]-1])
            assert t in VALID_TYPES, \
                f"Invalid type='{t}'. Must be one of {sorted(VALID_TYPES)}"

    def test_all_three_discrepancy_types_found(self):
        """Must identify all 3 discrepancy types."""
        wb = load_workbook(OUTPUT_FILE, data_only=True)
        ws = wb["Discrepancy Report"]
        cols = header_map(ws)
        found = {str(r[cols["type"]-1]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
        assert "AMOUNT_MISMATCH"     in found, "Missing AMOUNT_MISMATCH discrepancy"
        assert "MISSING_FROM_LEDGER" in found, "Missing MISSING_FROM_LEDGER discrepancy"
        assert "PHANTOM_ENTRY"       in found, "Missing PHANTOM_ENTRY discrepancy"
